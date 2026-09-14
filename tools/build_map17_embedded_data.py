#!/usr/bin/env python3
"""Build Map 17's authoritative snapshots at build time and embed them in the HTML.

Why: Census/USDM/USDA/CDC source servers are authoritative but do not all permit
browser CORS from GitHub Pages. This script runs server-side (GitHub Actions),
normalizes the data, and writes it into the standalone HTML so the public browser
never fetches those government data sources directly.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import sys
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

TARGET = Path("daily-maps/17-agricultural-drought-farm-exposure.html")
UA = "MitchellCo-DailyMap/1.0 (+https://christophermitchell012.github.io/wildfire-map/)"
STATES = "AL,AK,AZ,AR,CA,CO,CT,DC,DE,FL,GA,HI,ID,IL,IN,IA,KS,KY,LA,ME,MD,MA,MI,MN,MS,MO,MT,NE,NV,NH,NJ,NM,NY,NC,ND,OH,OK,OR,PA,RI,SC,SD,TN,TX,UT,VT,VA,WA,WV,WI,WY"
CENSUS_ZIP = "https://www2.census.gov/geo/docs/maps-data/data/gazetteer/2025_Gazetteer/2025_Gaz_counties_national.zip"
USDM = "https://usdmdataservices.unl.edu/api/CountyStatistics/GetDroughtSeverityStatisticsByAreaPercent"
USDA_PAGE = "https://www.nass.usda.gov/Publications/AgCensus/2022/Online_Resources/Ag_Census_Web_Maps/Data_download/"
CDC_PAGE = "https://www.atsdr.cdc.gov/place-health/php/svi/svi-data-documentation-download.html"


def get(url: str, timeout: int = 45, accept: str = "*/*") -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": accept})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def norm_fips(v) -> str | None:
    s = re.sub(r"\D", "", str(v or ""))
    if not s:
        return None
    return s.zfill(5)[:5]


def hrefs(page_url: str) -> list[str]:
    html = get(page_url).decode("utf-8", "replace")
    vals = re.findall(r'href=["\']([^"\']+)["\']', html, flags=re.I)
    return [urllib.parse.urljoin(page_url, x) for x in vals]


def build_centroids() -> dict:
    raw = get(CENSUS_ZIP)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        name = next(n for n in z.namelist() if re.search(r"counties.*\.txt$", n, re.I))
        text = z.read(name).decode("utf-8-sig", "replace")
    reader = csv.DictReader(io.StringIO(text), delimiter="|")
    out = {}
    for row in reader:
        fips = norm_fips(row.get("GEOID"))
        try:
            lat = float(row.get("INTPTLAT", ""))
            lon = float(row.get("INTPTLONG", ""))
        except ValueError:
            continue
        if fips:
            out[fips] = {
                "fips": fips,
                "name": row.get("NAME") or fips,
                "state": row.get("USPS") or "",
                "lat": round(lat, 6),
                "lon": round(lon, 6),
            }
    if len(out) < 3000:
        raise RuntimeError(f"Census centroid coverage too small: {len(out)}")
    return out


def build_usdm() -> tuple[dict, str]:
    # USDM is weekly. Ask for a short range and keep the newest record per county.
    # The official service documents CSV as its default response, so parse CSV
    # instead of assuming JSON from an Accept negotiation.
    end = datetime.now(timezone.utc).date() - timedelta(days=1)
    start = end - timedelta(days=15)
    fmt = lambda d: f"{d.month}/{d.day}/{d.year}"
    q = urllib.parse.urlencode({
        "aoi": STATES,
        "startdate": fmt(start),
        "enddate": fmt(end),
        "statisticsType": "1",
    })
    raw_text = get(f"{USDM}?{q}", accept="text/csv").decode("utf-8-sig", "replace")
    rows = list(csv.DictReader(io.StringIO(raw_text)))
    if not rows:
        raise RuntimeError(f"USDM returned no CSV records; prefix={raw_text[:200]!r}")
    out = {}
    dates = []
    for raw in rows:
        low = {str(k).lower(): v for k, v in raw.items()}
        fips = norm_fips(low.get("fips") or low.get("countyfips") or low.get("geoid"))
        if not fips:
            continue
        row = {
            "fips": fips,
            "name": low.get("name") or low.get("county") or low.get("countyname") or fips,
            "state": low.get("stateabbreviation") or low.get("state") or "",
            "date": low.get("mapdate") or low.get("date") or low.get("validdate") or "",
            "None": low.get("none"),
            "D0": low.get("d0"), "D1": low.get("d1"), "D2": low.get("d2"),
            "D3": low.get("d3"), "D4": low.get("d4"), "DSCI": low.get("dsci"),
        }
        if row["date"]:
            dates.append(str(row["date"]))
        prev = out.get(fips)
        if prev is None or str(row["date"]) > str(prev.get("date", "")):
            out[fips] = row
    if len(out) < 3000:
        sample_headers = list(rows[0].keys()) if rows else []
        raise RuntimeError(f"USDM county coverage too small: {len(out)}; headers={sample_headers}")
    return out, max(dates) if dates else "latest returned"


def build_usda() -> tuple[dict, str]:
    """Best-effort: discover the official NASS XLSX and extract crop-sales exposure."""
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError(f"openpyxl unavailable: {e}")

    candidates = [u for u in hrefs(USDA_PAGE) if re.search(r"\.xlsx?(?:$|\?)", u, re.I)]
    candidates = [u for u in candidates if "arcgis" not in u.lower() and "esri" not in u.lower()]
    if not candidates:
        raise RuntimeError("No first-party USDA XLSX link discovered")
    last = None
    for url in candidates:
        try:
            wb = load_workbook(io.BytesIO(get(url)), read_only=True, data_only=True)
            lookup = next((wb[s] for s in wb.sheetnames if "lookup" in s.lower()), None)
            if lookup is None:
                continue
            map_id = None
            title = None
            for row in lookup.iter_rows(values_only=True):
                vals = [str(v or "") for v in row]
                joined = " | ".join(vals).lower()
                if "crop" in joined and ("sold" in joined or "sales" in joined) and ("market value" in joined or "dollar" in joined):
                    for v in vals:
                        if re.fullmatch(r"y22_M\d+", v):
                            map_id = v
                            title = " | ".join(x for x in vals if x)
                            break
                if map_id:
                    break
            if not map_id:
                raise RuntimeError("Could not identify a crop-sales MapID")

            exposure = {}
            raw_values = {}
            for s in wb.sheetnames:
                ws = wb[s]
                rows = ws.iter_rows(values_only=True)
                try:
                    header = [str(x or "") for x in next(rows)]
                except StopIteration:
                    continue
                h = {name: i for i, name in enumerate(header)}
                fcol = next((h[x] for x in ("FIPSTEXT", "FIPS") if x in h), None)
                vname = f"{map_id}_valueNumeric"
                if fcol is None or vname not in h:
                    continue
                vcol = h[vname]
                for row in rows:
                    fips = norm_fips(row[fcol] if fcol < len(row) else None)
                    val = row[vcol] if vcol < len(row) else None
                    try:
                        val = float(val)
                    except (TypeError, ValueError):
                        continue
                    if fips and math.isfinite(val) and val >= 0:
                        raw_values[fips] = val
                break
            if not raw_values:
                raise RuntimeError("USDA crop-sales variable had no numeric county values")
            maxv = max(raw_values.values())
            den = math.log1p(maxv) or 1.0
            for fips, val in raw_values.items():
                exposure[fips] = round(math.log1p(val) / den, 6)
            return exposure, f"2022 Census of Agriculture; {map_id}; {title or 'crop sales'}"
        except Exception as e:
            last = e
    raise RuntimeError(f"USDA extraction failed: {last}")


def build_cdc_svi() -> tuple[dict, str]:
    """Best-effort: discover first-party CDC/ATSDR 2022 county CSV/ZIP."""
    links = hrefs(CDC_PAGE)
    cands = [u for u in links if "2022" in u.lower() and "county" in u.lower() and re.search(r"\.(csv|zip)(?:$|\?)", u, re.I)]
    cands = [u for u in cands if "arcgis" not in u.lower() and "esri" not in u.lower()]
    last = None
    for url in cands:
        try:
            b = get(url)
            if re.search(r"\.zip(?:$|\?)", url, re.I):
                with zipfile.ZipFile(io.BytesIO(b)) as z:
                    csv_name = next(n for n in z.namelist() if n.lower().endswith(".csv") and "county" in n.lower())
                    text = z.read(csv_name).decode("utf-8-sig", "replace")
            else:
                text = b.decode("utf-8-sig", "replace")
            reader = csv.DictReader(io.StringIO(text))
            out = {}
            for row in reader:
                fips = norm_fips(row.get("FIPS") or row.get("STCNTY"))
                try:
                    val = float(row.get("RPL_THEMES", ""))
                except (TypeError, ValueError):
                    continue
                if fips and val >= 0 and val != -999:
                    out[fips] = round(max(0.0, min(1.0, val)), 6)
            if len(out) < 3000:
                raise RuntimeError(f"SVI county coverage too small: {len(out)}")
            return out, "CDC/ATSDR SVI 2022 county overall percentile"
        except Exception as e:
            last = e
    raise RuntimeError(f"CDC SVI extraction failed: {last or 'no first-party CSV/ZIP link discovered'}")


def jsobj(obj) -> str:
    return json.dumps(obj, separators=(",", ":"), ensure_ascii=False)


def patch_html(centroids, usdm, usdm_date, usda, usda_meta, svi, svi_meta):
    text = TARGET.read_text(encoding="utf-8")
    generated = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    block = f'''// BEGIN MAP17 EMBEDDED DATA — generated by tools/build_map17_embedded_data.py\n        const EMBEDDED_META = Object.freeze({jsobj({"generatedUtc": generated, "usdmValid": usdm_date, "usda": usda_meta, "svi": svi_meta})});\n        const COUNTY_CENTROIDS = Object.freeze({jsobj(centroids)});\n        const USDM_SNAPSHOT = Object.freeze({jsobj(usdm)});\n        const USDA_AG_EXPOSURE = Object.freeze({jsobj(usda)});\n        const CDC_SVI = Object.freeze({jsobj(svi)});\n        // END MAP17 EMBEDDED DATA'''

    embedded_re = re.compile(r"// BEGIN MAP17 EMBEDDED DATA.*?// END MAP17 EMBEDDED DATA", re.S)
    if embedded_re.search(text):
        text = embedded_re.sub(block, text, count=1)
    else:
        constants_re = re.compile(
            r'\s*const STATES\s*=.*?const CDC_SVI\s*=\s*Object\.freeze\(\{\}\);',
            re.S,
        )
        if not constants_re.search(text):
            raise RuntimeError("Could not locate Map 17 source-constant block")
        text = constants_re.sub("\n        " + block, text, count=1)

    load_centroids = '''async function loadCentroids(run) {\n          if (run !== state.run) return false;\n          state.counties = new Map(Object.entries(COUNTY_CENTROIDS));\n          $("centroidCount").textContent = state.counties.size.toLocaleString();\n          setHealth("hCentroids", "Cached", "cached", `Embedded Census Gazetteer snapshot • ${EMBEDDED_META.generatedUtc}`);\n          return true;\n        }'''
    text, n1 = re.subn(
        r'async function loadCentroids\(run\) \{.*?\n\s*\}\n\s*async function loadDrought\(run\) \{',
        load_centroids + '\n        async function loadDrought(run) {',
        text,
        count=1,
        flags=re.S,
    )
    if n1 != 1:
        raise RuntimeError("Could not patch loadCentroids")

    load_drought_body = '''async function loadDrought(run) {\n          if (run !== state.run) return false;\n          state.drought = new Map(Object.entries(USDM_SNAPSHOT));\n          $("droughtCount").textContent = state.drought.size.toLocaleString();\n          setHealth("hDrought", "Cached", "cached", `Embedded USDM snapshot • valid ${EMBEDDED_META.usdmValid}`);\n          return true;\n        }\n        function initializeSnapshotHealth() {'''
    text, n2 = re.subn(
        r'async function loadDrought\(run\) \{.*?\n\s*\}\n\s*function initializeSnapshotHealth\(\) \{',
        load_drought_body,
        text,
        count=1,
        flags=re.S,
    )
    if n2 != 1:
        raise RuntimeError("Could not patch loadDrought")

    text = re.sub(r'\s*<script\s+src="https://cdnjs\.cloudflare\.com/ajax/libs/fflate/[^\"]+"\s*></script>', '', text)
    text = re.sub(r'\s*"services[235]\.arcgis\.com":\s*"ArcGIS",', '', text)
    text = re.sub(r'\s*"onemap\.cdc\.gov":\s*"CDC/ATSDR",', '', text)
    text = text.replace('"U.S. Drought Monitor REST"', '"U.S. Drought Monitor snapshot"')
    text = text.replace('"Census Gazetteer coordinates"', '"Census Gazetteer snapshot"')
    text = text.replace('"USDA Ag Census exposure"', '"USDA Ag Census snapshot"')
    text = text.replace('"CDC/ATSDR SVI"', '"CDC/ATSDR SVI snapshot"')
    TARGET.write_text(text, encoding="utf-8")


def main():
    print("Building Census county centroids…")
    centroids = build_centroids()
    print(f"  {len(centroids)} counties")
    print("Building current USDM county snapshot…")
    usdm, usdm_date = build_usdm()
    print(f"  {len(usdm)} counties; valid={usdm_date}")

    print("Building USDA crop-exposure snapshot…")
    try:
        usda, usda_meta = build_usda()
        print(f"  {len(usda)} counties")
    except Exception as e:
        print(f"WARNING USDA optional source unavailable: {e}", file=sys.stderr)
        usda, usda_meta = {}, f"Unavailable at build: {e}"

    print("Building CDC SVI snapshot…")
    try:
        svi, svi_meta = build_cdc_svi()
        print(f"  {len(svi)} counties")
    except Exception as e:
        print(f"WARNING CDC optional source unavailable: {e}", file=sys.stderr)
        svi, svi_meta = {}, f"Unavailable at build: {e}"

    patch_html(centroids, usdm, usdm_date, usda, usda_meta, svi, svi_meta)
    print(f"Patched {TARGET}")


if __name__ == "__main__":
    main()
